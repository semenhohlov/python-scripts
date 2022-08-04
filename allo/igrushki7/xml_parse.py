from xml.dom.minidom import parse, getDOMImplementation
import openpyxl
import os
import re
import datetime

def rename_old_price():
    if os.path.exists('export_i7.xml'):
        os.rename('export_i7.xml', 'prev.xml')
        print('Переименовую предыдущий export_i7.xml в prev.xml.')

def get_dom(name):
    dom = False
    try:
        dom = parse(name)
    except:
        print('Ошибка чтения файла: ', name)
    return dom

def overprice(dom, base):
    print('Переоценка...')
    impl = getDOMImplementation()
    over_pice = 2.20
    discount = 0.75
    newdoc = impl.createDocument(None, "offers", None)
    offers = newdoc.documentElement
    itemList = dom.getElementsByTagName('offer')
    total = len(itemList)
    percent = int(total / 100)
    current = 1
    tmp = 0
    for item in itemList:
        price =  item.getElementsByTagName('price')[0]
        value = float(price.childNodes[0].nodeValue)
        if value < 300:
            tmp += 1
            continue
        value = round(value * over_pice)
        # pos_name
        nameItem = item.getElementsByTagName('name')[0]
        name = str(nameItem.childNodes[0].nodeValue).replace('"', '')
        name = name.replace('р/у', 'на радиоуправлении')
        name = name.replace('р.у.', 'на радиоуправлении')
        name = name.replace('&qout;', '')
        name = name.replace('&amp;qout;', '')
        name = re.sub("\(\d*шт\)", "", name)
        nameItem.childNodes[0].nodeValue = name
        #description
        id = item.getAttribute('id')
        desc = base[id]
        descItem = item.getElementsByTagName('description')[0]
        # print(id, descItem.childNodes[0].nodeValue)
        try:
            descItem.childNodes[0].nodeValue = desc + ' ' + descItem.childNodes[0].nodeValue
        except:
            pass
        # if descItem.childNodes[0].nodeValue:
        #     descItem.childNodes[0].nodeValue = desc + ' ' + descItem.childNodes[0].nodeValue
        # else:
        #     newdoc = impl.createDocument(None, "description", None)
        #     newDesc = newdoc.documentElement
        #     newDesc.appendChild(newdoc.createTextNode(desc))
        #     item.appendChild(newDesc)


        newdoc = impl.createDocument(None, "oldprice", None)
        oldprice = newdoc.documentElement
        oldvalue = newdoc.createTextNode(str(value))
        oldprice.appendChild(oldvalue)
        price.childNodes[0].nodeValue = round(value * discount)
        item.appendChild(oldprice)
        offers.appendChild(item)
        tmp += 1
        if tmp > percent:
            tmp = 0
            current += 1
            print('\r{}%'.format(current), end='')
    print()
    return offers

def save_export(dom):
    f = open('export_i7.xml', 'w')
    f.write('<?xml version="1.0" encoding="UTF-8"?>\n')
    dom.writexml(f)
    f.close()
    print('Файл export_i7.xml готов для загрузки.')

def append_na_items(export, prev):
    print('Нет в наличии...')
    new_items = []
    na_items = 0
    itemsList = export.getElementsByTagName('offers')[0]
    itemList = export.getElementsByTagName('offer')
    for item in itemList:
        new_items.append(item.getAttribute('id'))
    itemList = prev.getElementsByTagName('offer')
    total = len(itemList)
    percent = int(total / 100)
    current = 1
    tmp = 0
    for item in itemList:
        try:
            new_items.index(item.getAttribute('id'))
        except:
            if item.getAttribute('available'):
                item.setAttribute('available', False)
                itemsList.appendChild(item)
                na_items += 1
        tmp += 1
        if tmp > percent:
            tmp = 0
            current += 1
            print('\r{}%'.format(current), end='')
    print()
    print("Нет в наличии {} позиций.".format(na_items))

def load_description(filename="desc.xlsx"):
    data = {}
    try:
        workbook = openpyxl.load_workbook(filename)
    except:
        return data
    worksheet = workbook.active
    for row in worksheet.rows:
        data[row[0].value] = row[1].value
    return data

def main():
    rename_old_price()
    export = get_dom('igrushki7.yml')
    prev = get_dom('prev.xml')
    if export == False:
        return
    base = load_description()
    newItems = overprice(export, base)

    impl = getDOMImplementation()
    newdoc = impl.createDocument(None, "yml_catalog", None)
    yml = newdoc.documentElement
    yml.setAttribute('date', datetime.datetime.now().strftime('{%Y-%m-%d %H:%M}'))
    newdoc = impl.createDocument(None, "shop", None)
    shop = newdoc.documentElement
    yml.appendChild(shop)
    newdoc = impl.createDocument(None, "name", None)
    name = newdoc.documentElement
    name.appendChild(newdoc.createTextNode('Игрушки 7'))
    newdoc = impl.createDocument(None, "company", None)
    company = newdoc.documentElement
    company.appendChild(newdoc.createTextNode('Игрушки 7'))
    newdoc = impl.createDocument(None, "url", None)
    url = newdoc.documentElement
    url.appendChild(newdoc.createTextNode('https://igrushki7.ua'))
    newdoc = impl.createDocument(None, "currencies", None)
    currencies = newdoc.documentElement
    newdoc = impl.createDocument(None, "currency", None)
    currency = newdoc.documentElement
    currency.setAttribute('id', 'UAH')
    currency.setAttribute('rate', '1')
    currencies.appendChild(currency)
    shop.appendChild(name)
    shop.appendChild(company)
    shop.appendChild(url)
    shop.appendChild(currencies)
    cats = export.getElementsByTagName('categories')[0]
    shop.appendChild(cats)
    shop.appendChild(newItems)

    if prev != False:
        append_na_items(yml, prev)
    save_export(yml)

if __name__ == '__main__':
    main()